#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_xlsx_h10_workload.py — (Hailo-10H, npu-rpi5) det/seg/pose 워크로드 조합 실험
결과를 컬럼설명/전체_21회/조건별_3회평균/그래프 4개 시트짜리 xlsx로 정리한다.

스케줄러 파라미터(priority/threshold/timeout)는 전혀 건드리지 않고 HailoRT 기본값
그대로 둔 채 워크로드 조합만 7가지로 바꾼 실험이다(8L의 default_workload와 동일 설계).

그래프 시트에는 엑셀 네이티브 차트 3개가 들어간다(이미지 붙여넣기 아님 — 엑셀에서
클릭하면 값이 살아 있는 진짜 차트 객체):
  차트 A: 조건별 모델별 평균 Latency (묶은 막대)
  차트 B: 조건별 처리량 FPS (누적 막대 — 쌓인 높이가 곧 합계 FPS)
  차트 C: 조건별 호스트 CPU 사용률 (막대)

사용법: python build_xlsx_h10_workload.py [21runs.csv] [avg.csv] [출력.xlsx]
        (인자 없으면 csv/ 안의 기본 경로를 쓴다)
"""
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
DEF_RAW = os.path.join(HERE, 'csv', 'results_det_seg_pose_workload_21runs.csv')
DEF_AVG = os.path.join(HERE, 'csv', 'results_det_seg_pose_workload_avg.csv')
DEF_OUT = os.path.join(HERE, 'csv', 'results_det_seg_pose_workload_combined.xlsx')

CONDITIONS = ['det', 'seg', 'pose', 'det_seg', 'det_pose', 'seg_pose', 'det_seg_pose']
MODELS = ['det', 'seg', 'pose']
# 모델 정체성 = 색상 (hailo_8L 실험 시리즈 카테고리 팔레트와 동일 계열)
MODEL_COLOR = {'det': '2A78D6', 'seg': 'EB6834', 'pose': '2CA02C'}
CPU_COLOR = '7F5FC4'

HDR_FILL = PatternFill('solid', fgColor='1F4E79')
HDR_FONT = Font(color='FFFFFF', bold=True)
SEC_FONT = Font(bold=True, size=12, color='1F4E79')

COLUMN_DESC = [
    ("실험 조건",
     "Hailo-10H(npu-rpi5, RPi5, hailort 5.3.0). det=yolov8s / seg=yolov8s_seg / "
     "pose=yolov8s_pose 3개 모델의 조합 7가지(단독 3 + 쌍 3 + 트리플 1)를 각 3회씩 "
     "= 21회 실행. **스케줄러 파라미터(priority/threshold/timeout)는 전혀 설정하지 "
     "않고 HailoRT 기본값 그대로 둔 채 워크로드 조합만 바꿨다.** 데이터셋은 8L 실험과 "
     "동일한 sampled_val2017 673장. 모델별로 독립 스레드가 하나씩 붙어 "
     "'전처리 → run_async → wait'를 673회 반복한다."),
    ("데이터 출처",
     "장비에서 받은 원본 csv/results_det_seg_pose_workload.csv는 스윕 3세트가 append로 "
     "누적된 58행이다. 그중 이 표에 쓴 21행은 세트 C(원본 38~58행)로, "
     "logs/<label>_run<N>.log 21개 파일과 값이 1:1로 대응돼 추적이 가능한 세트다. "
     "세트 A(원본 1~21행, logs/sweep_full.log)도 완주했고 세트 C와 1% 이내로 일치한다 "
     "— results_det_seg_pose_workload_setA_21runs.csv 참고. 세트 B(원본 22~37행)는 "
     "seg_pose run1에서 중단된 16행이라 제외했다."),
    ("label", "워크로드 조합. det / seg / pose / det_seg / det_pose / seg_pose / det_seg_pose"),
    ("run_id", "같은 조건의 반복 순번(1~3)"),
    ("use_det / use_seg / use_pose", "해당 실행에서 그 모델이 활성인지 여부(1=활성, 0=비활성)"),
    ("<model>_frame_count", "모델별로 정상 처리된 프레임 수(전 조건 673 = 데이터셋 전량)"),
    ("<model>_avg_latency_ms",
     "장당 평균 추론 시간(ms) — run_async() 제출 ~ wait() 완료 왕복. **전처리는 빠져 "
     "있다.** 비활성 모델은 빈칸."),
    ("<model>_fps",
     "모델별 처리량 = 프레임 수 ÷ 그 모델 워커의 전체 루프 시간. **전처리가 포함된다.** "
     "그래서 fps != 1000/avg_latency_ms 이다(전처리가 latency에는 빠지고 FPS에는 들어감). "
     "예: 3모델 동시 조건에서 prep 4.8ms + infer 25.9ms = 30.7ms -> 약 32.6 FPS."),
    ("total_fps (평균 시트/그래프에만)",
     "활성 모델들의 fps 합 = 그 조건에서 디바이스가 낸 전체 처리량. 조합 간 성능을 비교할 "
     "때 가장 안전한 지표다."),
    ("cpu_percent / mem_percent",
     "실행 중 **호스트(RPi5) 시스템 전체** 사용률(%). /proc/stat, /proc/meminfo를 300ms "
     "간격으로 샘플링해 평균낸 값이며 디바이스(NPU) 사용률이 아니다."),
    ("n_runs (평균 시트에만)", "그룹에 포함된 반복 횟수(=3)"),
    ("이 실험에서 얻을 수 없는 것",
     "H10은 VDevice/스케줄러가 디바이스 내부 hailort_server 프로세스에 있고 호스트 "
     "libhailort는 RPC 클라이언트라, 호스트에서 만드는 .hrtt 트레이스가 항상 0바이트다"
     "(HRTT_ON_HAILO10H.md 참고). 따라서 8L 실험에 있던 npu_percent / switches_per_s / "
     "idle_time_pct / activation_ms 같은 스케줄러 내부 지표가 이 표에는 없다. "
     "이 스윕(2026-08-26)은 host_trace 구현(2026-08-27)보다 하루 앞서 돌아서 프레임 단위 "
     "타임라인(traces/)도 남아 있지 않다."),
]


def read_dicts(path):
    with open(path, newline='', encoding='utf-8') as f:
        return list(csv.DictReader(f))


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def style_header(ws, row=1):
    for c in ws[row]:
        if c.value is not None:
            c.fill = HDR_FILL
            c.font = HDR_FONT
            c.alignment = Alignment(horizontal='center', vertical='center')


def autosize(ws, min_w=10, max_w=26):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[letter].width = min(max(longest + 2, min_w), max_w)


def sheet_desc(wb):
    ws = wb.create_sheet('컬럼설명')
    ws.append(['항목', '설명'])
    style_header(ws)
    for k, v in COLUMN_DESC:
        ws.append([k, v])
    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 110
    for r in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        r[0].alignment = Alignment(wrap_text=True, vertical='top')
    for r in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        r[0].alignment = Alignment(vertical='top')
        r[0].font = Font(bold=True)
    ws.freeze_panes = 'A2'
    return ws


def sheet_raw(wb, raw):
    ws = wb.create_sheet('전체_21회')
    hdr = list(raw[0].keys())
    ws.append(hdr)
    style_header(ws)
    for r in raw:
        ws.append([num(r[h]) if num(r[h]) is not None else (r[h] or None) for h in hdr])
    ws.freeze_panes = 'C2'
    autosize(ws)
    return ws


def build_avg(raw):
    """조건별 3회 평균을 직접 계산한다(avg.csv에 의존하지 않고 21행에서 바로 낸다)."""
    out = []
    for cond in CONDITIONS:
        rs = [r for r in raw if r['label'] == cond]
        if not rs:
            continue
        row = {'label': cond, 'n_runs': len(rs)}
        total = 0.0
        for m in MODELS:
            lat = [num(r['%s_avg_latency_ms' % m]) for r in rs]
            fps = [num(r['%s_fps' % m]) for r in rs]
            lat = [v for v in lat if v is not None]
            fps = [v for v in fps if v is not None]
            row['%s_avg_latency_ms' % m] = round(sum(lat) / len(lat), 4) if lat else None
            row['%s_fps' % m] = round(sum(fps) / len(fps), 4) if fps else None
            if fps:
                total += sum(fps) / len(fps)
        row['total_fps'] = round(total, 4)
        for c in ('cpu_percent', 'mem_percent'):
            vs = [num(r[c]) for r in rs]
            vs = [v for v in vs if v is not None]
            row[c] = round(sum(vs) / len(vs), 4) if vs else None
        out.append(row)
    return out


def sheet_avg(wb, avg):
    ws = wb.create_sheet('조건별_3회평균')
    hdr = (['label', 'n_runs']
           + ['%s_avg_latency_ms' % m for m in MODELS]
           + ['%s_fps' % m for m in MODELS]
           + ['total_fps', 'cpu_percent', 'mem_percent'])
    ws.append(hdr)
    style_header(ws)
    for r in avg:
        ws.append([r.get(h) for h in hdr])
    ws.freeze_panes = 'B2'
    autosize(ws, min_w=12)
    return ws


def write_block(ws, top, title, header, rows):
    """그래프 시트에 차트용 데이터 블록을 쓰고 (헤더행, 마지막행)을 돌려준다."""
    ws.cell(row=top, column=1, value=title).font = SEC_FONT
    hdr_row = top + 1
    for j, h in enumerate(header, start=1):
        c = ws.cell(row=hdr_row, column=j, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(horizontal='center')
    for i, row in enumerate(rows, start=hdr_row + 1):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)
    return hdr_row, hdr_row + len(rows)


def color_series(chart, colors):
    for s, hexc in zip(chart.series, colors):
        s.graphicalProperties.solidFill = hexc
        s.graphicalProperties.line.noFill = True


def sheet_chart(wb, avg):
    ws = wb.create_sheet('그래프')
    labels = [r['label'] for r in avg]

    # --- 차트 A: 조건별 모델별 평균 Latency (묶은 막대) ---
    rows_a = [[r['label']] + [r.get('%s_avg_latency_ms' % m) for m in MODELS] for r in avg]
    hdr_a, last_a = write_block(ws, 1, '차트 A 데이터 — 조건별 모델별 평균 Latency (ms)',
                                ['조건'] + MODELS, rows_a)
    ch_a = BarChart()
    ch_a.type = 'col'
    ch_a.grouping = 'clustered'
    ch_a.title = 'Hailo-10H 조건별 평균 Latency (ms, 3회 평균)'
    ch_a.y_axis.title = 'avg_latency_ms'
    ch_a.x_axis.title = '워크로드 조합'
    ch_a.height, ch_a.width = 9, 20
    ch_a.add_data(Reference(ws, min_col=2, max_col=1 + len(MODELS),
                            min_row=hdr_a, max_row=last_a), titles_from_data=True)
    ch_a.set_categories(Reference(ws, min_col=1, max_col=1, min_row=hdr_a + 1, max_row=last_a))
    color_series(ch_a, [MODEL_COLOR[m] for m in MODELS])
    ws.add_chart(ch_a, 'G2')

    # --- 차트 B: 조건별 처리량 FPS (누적 막대 = 쌓인 높이가 합계 FPS) ---
    top_b = last_a + 2
    rows_b = [[r['label']] + [r.get('%s_fps' % m) for m in MODELS] + [r['total_fps']]
              for r in avg]
    hdr_b, last_b = write_block(ws, top_b, '차트 B 데이터 — 조건별 처리량 FPS (합계 = 누적 높이)',
                                ['조건'] + MODELS + ['합계FPS(참고)'], rows_b)
    ch_b = BarChart()
    ch_b.type = 'col'
    ch_b.grouping = 'stacked'
    ch_b.overlap = 100
    ch_b.title = 'Hailo-10H 조건별 처리량 (FPS, 모델별 누적 = 합계)'
    ch_b.y_axis.title = 'FPS'
    ch_b.x_axis.title = '워크로드 조합'
    ch_b.height, ch_b.width = 9, 20
    ch_b.add_data(Reference(ws, min_col=2, max_col=1 + len(MODELS),
                            min_row=hdr_b, max_row=last_b), titles_from_data=True)
    ch_b.set_categories(Reference(ws, min_col=1, max_col=1, min_row=hdr_b + 1, max_row=last_b))
    color_series(ch_b, [MODEL_COLOR[m] for m in MODELS])
    ch_b.dLbls = DataLabelList()
    ch_b.dLbls.showVal = True
    ws.add_chart(ch_b, 'G21')

    # --- 차트 C: 조건별 호스트 CPU 사용률 ---
    top_c = last_b + 2
    rows_c = [[r['label'], r['cpu_percent']] for r in avg]
    hdr_c, last_c = write_block(ws, top_c, '차트 C 데이터 — 조건별 호스트 CPU 사용률 (%)',
                                ['조건', 'cpu_percent'], rows_c)
    ch_c = BarChart()
    ch_c.type = 'col'
    ch_c.title = '호스트(RPi5) CPU 사용률 (%, 3회 평균)'
    ch_c.y_axis.title = 'cpu_percent'
    ch_c.x_axis.title = '워크로드 조합'
    ch_c.height, ch_c.width = 9, 20
    ch_c.add_data(Reference(ws, min_col=2, max_col=2, min_row=hdr_c, max_row=last_c),
                  titles_from_data=True)
    ch_c.set_categories(Reference(ws, min_col=1, max_col=1, min_row=hdr_c + 1, max_row=last_c))
    color_series(ch_c, [CPU_COLOR])
    ch_c.legend = None
    ws.add_chart(ch_c, 'G40')

    ws.column_dimensions['A'].width = 18
    for j in range(2, 6):
        ws.column_dimensions[get_column_letter(j)].width = 14
    assert len(labels) == len(CONDITIONS)
    return ws


def main():
    raw_p = sys.argv[1] if len(sys.argv) > 1 else DEF_RAW
    avg_p = sys.argv[2] if len(sys.argv) > 2 else DEF_AVG
    out_p = sys.argv[3] if len(sys.argv) > 3 else DEF_OUT

    raw = read_dicts(raw_p)
    if len(raw) != 21:
        print('경고: 원본이 21행이 아니다 (%d행) — %s' % (len(raw), raw_p))
    avg = build_avg(raw)

    wb = Workbook()
    wb.remove(wb.active)
    sheet_desc(wb)
    sheet_raw(wb, raw)
    sheet_avg(wb, avg)
    sheet_chart(wb, avg)
    wb.save(out_p)

    print('시트: %s' % ', '.join(wb.sheetnames))
    print('xlsx -> %s' % out_p)
    print('(참고) avg.csv 경로는 인자로만 받고 평균은 21행에서 직접 계산함: %s'
          % os.path.basename(avg_p))


if __name__ == '__main__':
    main()
